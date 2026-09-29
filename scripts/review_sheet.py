"""Sharia review of the content drafts as one spreadsheet (M1 sign-off).

    PYTHONPATH=. python scripts/review_sheet.py export   → docs/review/Nun_content_review.xlsx
    PYTHONPATH=. python scripts/review_sheet.py import   ← the same file, filled in by the reviewer

export: one sheet each for Names, Dhikr, Inscriptions, plus the open Questions (content + sources), each row with a
Decision drop-down (Approve / Reject / Edit) and notes columns. Nothing in it is authoritative until approved.
import: records every decision in data/lists/review_decisions.json, writes the APPROVED items (with reviewed_by) to
data/lists/{names.jsonl, dhikr.jsonl, inscriptions.yaml}, lists "Edit" rows for follow-up, and appends a
sign-off entry to docs/REVIEW_LOG.md. Rows left without a decision stay pending.
"""

from __future__ import annotations

import hashlib
import json
import sys
from datetime import date
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

LISTS = Path("data/lists")
SHEET = Path("docs/review/Nun_content_review.xlsx")
DECISIONS = LISTS / "review_decisions.json"
REVIEW_LOG = Path("docs/REVIEW_LOG.md")
CHOICES = ["Approve", "Reject", "Edit"]
HEAD = PatternFill("solid", fgColor="1F5E57")
DECIDE = PatternFill("solid", fgColor="FFF4D6")

QUESTIONS = [
    (
        "q1-names-list",
        "The 99 Names: the enumeration is the list appended to the hadith in at-Tirmidhi (3507); the "
        "hadith itself in al-Bukhari and Muslim has no list. Present it as 'the commonly known list', or use another?",
    ),
    (
        "q2-verse-or-dhikr",
        "Phrases that are also Quran text (e.g. «حسبنا الله ونعم الوكيل» 3:173, «ما شاء الله»): "
        "when photographed alone, show them as a verse, as a dhikr, or both?",
    ),
    (
        "q3-ambiguous",
        "A fragment found in several verses (e.g. «إن الله على كل شيء قدير», 11 places): the app "
        "plans to show ALL references rather than pick one. Agree?",
    ),
    (
        "q4-more-verses",
        "Coverage: 184 inscription spans exist; the plan targets ~500. Which further verses commonly "
        "inscribed in mosques/museums should be added? (list references)",
    ),
    (
        "s1-french",
        "French translation: french_montada (Noor International, 1.0.0) or french_rashid (Rachid Maach, "
        "1.0.3)? Is the chosen one a King Fahd Complex edition or listed on quranpedia.net?",
    ),
    ("s2-ur-id-zh", "Confirm on quranpedia.net: urdu_junagarhi 1.1.3, indonesian_complex 1.0.1, chinese_makin 1.0.2."),
    (
        "s3-tafsir",
        "Tafsir: al-Muyassar (King Fahd Complex; not named in the package), Dorar's «موسوعة التفسير», or "
        "both (Dorar first where available)?",
    ),
]


def span_id(e: dict) -> str:
    """Stable id of an inscription span: several fragments can share one reference (e.g. parts of 2:255)."""
    digest = hashlib.sha1((e["fragment"] or "").encode("utf-8")).hexdigest()[:6]
    return f"{e['ref']}#{digest}"


def _jsonl(p: Path) -> list[dict]:
    return [json.loads(line) for line in p.open(encoding="utf-8")]


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list], widths: list[int], decide_col: int):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[decide_col - 1].fill = DECIDE
    dv = DataValidation(type="list", formula1='"' + ",".join(CHOICES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    col = ws.cell(1, decide_col).column_letter
    dv.add(f"{col}2:{col}{ws.max_row}")
    ws.freeze_panes = "B2"
    return ws


def export() -> None:
    names, dhikr = _jsonl(LISTS / "names.draft.jsonl"), _jsonl(LISTS / "dhikr.draft.jsonl")
    ins = yaml.safe_load((LISTS / "inscriptions.draft.yaml").read_text(encoding="utf-8"))["spans"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    for line in [
        ["Nūn · content review (M1)"],
        ["Reviewer name:", ""],
        [],
        ["For every row on the Names / Dhikr / Inscriptions sheets, choose a Decision (yellow column):"],
        ["  Approve = correct and may be shown to users · Reject = remove · Edit = keep with the correction you write"],
        ["Answer each item on the Questions sheet."],
        ["Machine checks (Quran occurrences, references) were computed against the verified corpus."],
        ["Hadith references marked 'UNVERIFIED' are from memory: confirm book, number and grading, or reject."],
        ["Rows without a decision stay pending and are NOT used by the app."],
        [],
        ["راجع كل صف واختر القرار (موافقة / رفض / تعديل)، وأجب عن الأسئلة. الصفوف دون قرار لا تُستخدم في التطبيق."],
    ]:
        ws.append(line)
    ws["A1"].font = Font(bold=True, size=14)
    ws["B2"].fill = DECIDE
    ws.column_dimensions["A"].width = 110

    _sheet(
        wb,
        "Names",
        [
            "id",
            "Name",
            "In the Quran as written (count · first refs)",
            "Without «ال» (plain word matches)",
            "Decision",
            "Correction (if Edit)",
            "Meaning (approved source)",
            "Reviewer notes",
        ],
        [
            [
                n["id"],
                n["text"],
                f"{n['quran_as_written']['count']} · {', '.join(n['quran_as_written']['refs'])}",
                (
                    f"{n['quran_without_al']['count']} · {', '.join(n['quran_without_al']['refs'])}"
                    if n["quran_without_al"]
                    else ""
                ),
                "",
                "",
                "",
                "",
            ]
            for n in names
        ],
        [9, 22, 30, 30, 12, 24, 40, 40],
        5,
    )
    _sheet(
        wb,
        "Dhikr",
        [
            "id",
            "Phrase",
            "Verbatim in the Quran at",
            "Hadith candidate (UNVERIFIED)",
            "dorar.net search",
            "Decision",
            "Correction (if Edit)",
            "Confirmed source: book · number · grading",
            "Reviewer notes",
        ],
        [
            [
                d["id"],
                d["text"],
                ", ".join(d["quran_verbatim_refs"][:8]) + (" …" if len(d["quran_verbatim_refs"]) > 8 else ""),
                (f"{d['hadith_candidate']['book']} {d['hadith_candidate']['number']}" if d["hadith_candidate"] else ""),
                d["dorar_search"],
                "",
                "",
                "",
                "",
            ]
            for d in dhikr
        ],
        [22, 34, 26, 22, 30, 12, 24, 36, 36],
        6,
    )
    _sheet(
        wb,
        "Inscriptions",
        [
            "id",
            "ref",
            "Inscribed fragment (corpus text)",
            "Found in",
            "Match",
            "Also occurs at",
            "Decision",
            "Reviewer notes",
        ],
        [
            [
                span_id(e),
                e["ref"],
                e["fragment"],
                ", ".join(f"{k}×{v}" for k, v in e["sources"].items()),
                e["match"],
                ", ".join(e["also_occurs_at"][:10]),
                "",
                "",
            ]
            for e in ins
        ],
        [16, 10, 60, 22, 14, 26, 12, 36],
        7,
    )
    q = wb.create_sheet("Questions")
    q.append(["id", "Question", "Answer", "Reviewer notes"])
    for c in q[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
    for qid, text in QUESTIONS:
        q.append([qid, text, "", ""])
    for i, w in enumerate([18, 90, 50, 40], 1):
        q.column_dimensions[q.cell(1, i).column_letter].width = w
    for row in q.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[2].fill = DECIDE
    SHEET.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SHEET)
    print(f"→ {SHEET}: {len(names)} names · {len(dhikr)} dhikr · {len(ins)} inscriptions · {len(QUESTIONS)} questions")


def _rows(ws, key_col: int, decide_col: int, extra: dict[str, int]) -> dict[str, dict]:
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[key_col - 1]:
            continue
        d = (row[decide_col - 1] or "").strip()
        out[str(row[key_col - 1])] = {"decision": d or None} | {k: row[i - 1] for k, i in extra.items()}
    return out


def do_import() -> None:
    wb = load_workbook(SHEET)
    reviewer = (wb["Read me"]["B2"].value or "").strip()
    if not reviewer:
        sys.exit("Fill in 'Reviewer name' on the Read me sheet first.")
    dec = {
        "reviewer": reviewer,
        "date": date.today().isoformat(),
        "names": _rows(wb["Names"], 1, 5, {"correction": 6, "meaning": 7, "notes": 8}),
        "dhikr": _rows(wb["Dhikr"], 1, 6, {"correction": 7, "source": 8, "notes": 9}),
        "inscriptions": _rows(wb["Inscriptions"], 1, 7, {"ref": 2, "notes": 8}),
        "questions": {
            r[0]: {"answer": r[2], "notes": r[3]}
            for r in wb["Questions"].iter_rows(min_row=2, values_only=True)
            if r[0]
        },
    }
    for kind in ("names", "dhikr", "inscriptions"):
        bad = {k: v["decision"] for k, v in dec[kind].items() if v["decision"] not in (None, *CHOICES)}
        if bad:
            sys.exit(f"{kind}: unknown decisions {bad}")
    DECISIONS.write_text(json.dumps(dec, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    def approved(kind: str, key: str) -> dict | None:
        d = dec[kind].get(key)
        return d if d and d["decision"] == "Approve" else None

    names = [
        n | {"status": "approved", "reviewed_by": reviewer}
        for n in _jsonl(LISTS / "names.draft.jsonl")
        if approved("names", n["id"])
    ]
    dhikr = []
    for d in _jsonl(LISTS / "dhikr.draft.jsonl"):
        if a := approved("dhikr", d["id"]):
            hc = d["hadith_candidate"]
            if hc and not a.get("source"):
                hc = None  # an unconfirmed hadith reference is never kept
            dhikr.append(
                d
                | {
                    "hadith_candidate": None,
                    "source": a.get("source") or (None if hc is None else hc),
                    "status": "approved",
                    "reviewed_by": reviewer,
                }
            )
    spans = yaml.safe_load((LISTS / "inscriptions.draft.yaml").read_text(encoding="utf-8"))["spans"]
    ins = [e | {"status": "approved", "reviewed_by": reviewer} for e in spans if approved("inscriptions", span_id(e))]
    (LISTS / "names.jsonl").write_text(
        "".join(json.dumps(x, ensure_ascii=False) + "\n" for x in names), encoding="utf-8"
    )
    (LISTS / "dhikr.jsonl").write_text(
        "".join(json.dumps(x, ensure_ascii=False) + "\n" for x in dhikr), encoding="utf-8"
    )
    (LISTS / "inscriptions.yaml").write_text(
        yaml.safe_dump({"spans": ins}, allow_unicode=True, sort_keys=False), encoding="utf-8"
    )

    def count(kind: str) -> dict:
        c: dict = {}
        for v in dec[kind].values():
            c[v["decision"] or "pending"] = c.get(v["decision"] or "pending", 0) + 1
        return c

    summary = {k: count(k) for k in ("names", "dhikr", "inscriptions")}
    answered = sum(1 for v in dec["questions"].values() if v["answer"])
    edits = {k: [i for i, v in dec[k].items() if v["decision"] == "Edit"] for k in ("names", "dhikr", "inscriptions")}
    with REVIEW_LOG.open("a", encoding="utf-8") as f:
        f.write(
            f"| {dec['date']} | M1 lists via `{SHEET.name}`: names {summary['names']}, dhikr {summary['dhikr']}, "
            f"inscriptions {summary['inscriptions']}; questions answered {answered}/{len(dec['questions'])} "
            f"(`{DECISIONS}`) | imported | {reviewer} |\n"
        )
    print(
        json.dumps(
            {"reviewer": reviewer, **summary, "questions_answered": answered, "edits_to_apply": edits},
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    {"export": export, "import": do_import}[sys.argv[1] if len(sys.argv) > 1 else "export"]()
