"""Review spreadsheet round trip: export → reviewer fills decisions → import keeps only approved rows."""

import json
import sys
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import review_sheet as rs  # noqa: E402

pytestmark = pytest.mark.skipif(not Path("data/lists/names.draft.jsonl").exists(), reason="run `make lists` first")


@pytest.fixture
def paths(tmp_path, monkeypatch):
    lists = tmp_path / "lists"
    lists.mkdir()
    for f in ("names.draft.jsonl", "dhikr.draft.jsonl", "inscriptions.draft.yaml"):
        (lists / f).write_bytes((Path("data/lists") / f).read_bytes())
    monkeypatch.setattr(rs, "LISTS", lists)
    monkeypatch.setattr(rs, "SHEET", tmp_path / "review.xlsx")
    monkeypatch.setattr(rs, "DECISIONS", lists / "review_decisions.json")
    monkeypatch.setattr(rs, "REVIEW_LOG", tmp_path / "REVIEW_LOG.md")
    return lists


def test_round_trip(paths) -> None:
    rs.export()
    wb = load_workbook(rs.SHEET)
    wb["Read me"]["B2"] = "Test Reviewer"
    wb["Names"]["E2"], wb["Names"]["E3"] = "Approve", "Reject"
    # approve a dhikr that has an unverified hadith candidate, WITHOUT confirming the source
    dh = wb["Dhikr"]
    row = next(r for r in range(2, dh.max_row + 1) if dh.cell(r, 4).value)
    dh.cell(row, 6).value = "Approve"
    wb["Inscriptions"]["G2"] = "Approve"
    wb["Questions"]["C2"] = "commonly known list"
    wb.save(rs.SHEET)

    rs.do_import()

    names = [json.loads(x) for x in (paths / "names.jsonl").read_text(encoding="utf-8").splitlines()]
    assert len(names) == 1 and names[0]["reviewed_by"] == "Test Reviewer"
    dhikr = [json.loads(x) for x in (paths / "dhikr.jsonl").read_text(encoding="utf-8").splitlines()]
    assert len(dhikr) == 1 and dhikr[0]["source"] is None and dhikr[0]["hadith_candidate"] is None
    ins = yaml.safe_load((paths / "inscriptions.yaml").read_text(encoding="utf-8"))["spans"]
    assert len(ins) == 1
    dec = json.loads(rs.DECISIONS.read_text(encoding="utf-8"))
    assert dec["names"]["n:allah"]["decision"] == "Approve" and dec["names"]["n:01"]["decision"] == "Reject"
    assert "Test Reviewer" in rs.REVIEW_LOG.read_text(encoding="utf-8")


def test_import_requires_reviewer_name(paths) -> None:
    rs.export()
    with pytest.raises(SystemExit):
        rs.do_import()
